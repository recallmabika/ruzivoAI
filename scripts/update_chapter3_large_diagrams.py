import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

SRC_DOCX = r"C:\Users\recal\Downloads\Ruzivo_Chapter3_Complete_Original_Backup.docx"
OUT_DOCX = r"C:\Users\recal\Downloads\Ruzivo_Chapter3_Complete.docx"

DIAGRAMS_DIR = r"c:\Users\recal\Desktop\Level 2.2 Project\ruzivo\docs\diagrams"
FIG_3_1 = os.path.join(DIAGRAMS_DIR, "figure_3_1_architecture.png")
FIG_3_2 = os.path.join(DIAGRAMS_DIR, "figure_3_2_activity_diagram.png")
FIG_3_3 = os.path.join(DIAGRAMS_DIR, "figure_3_3_rag_pipeline.png")
FIG_3_4 = os.path.join(DIAGRAMS_DIR, "figure_3_4_authentication.png")

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>\n'
                        f'  <w:top w:val="{"single" if top else "none"}" w:sz="{sz}" w:space="0" w:color="{top if top else "auto"}"/>\n'
                        f'  <w:left w:val="{"single" if left else "none"}" w:sz="{sz}" w:space="0" w:color="{left if left else "auto"}"/>\n'
                        f'  <w:bottom w:val="{"single" if bottom else "none"}" w:sz="{sz}" w:space="0" w:color="{bottom if bottom else "auto"}"/>\n'
                        f'  <w:right w:val="{"single" if right else "none"}" w:sz="{sz}" w:space="0" w:color="{right if right else "auto"}"/>\n'
                        f'</w:tcBorders>')
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>\n'
                      f'  <w:top w:w="{top}" w:type="dxa"/>\n'
                      f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
                      f'  <w:left w:w="{left}" w:type="dxa"/>\n'
                      f'  <w:right w:w="{right}" w:type="dxa"/>\n'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.5):
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    for r in p.runs:
        r.font.name = font_name
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.italic = italic

def make_caption_p(p, text):
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.italic = False

def update_chapter_3_with_large_diagrams():
    doc = docx.Document(SRC_DOCX)
    print(f"Loaded {SRC_DOCX} ({len(doc.paragraphs)} paras, {len(doc.tables)} tables)")

    # 1. Grab XML pointers to existing tables BEFORE doing any removals
    tbl_hw_elem = doc.tables[0]._element
    tbl_sw_elem = doc.tables[1]._element
    tbl2_elem = doc.tables[2]._element
    tbls_fig32_elem = [doc.tables[i]._element for i in range(5, 15)]
    tbls_fig33_elem = [doc.tables[i]._element for i in range(15, 23)]
    tbls_fig34_elem = [doc.tables[i]._element for i in range(23, 34)]

    # 2. Insert Figure 3.1 as a large, wide, high-resolution picture (6.5 inches)
    p_img31 = doc.add_paragraph()
    p_img31.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img31.paragraph_format.space_before = Pt(10)
    p_img31.paragraph_format.space_after = Pt(4)
    r31 = p_img31.add_run()
    r31.add_picture(FIG_3_1, width=Inches(6.5))
    tbl2_elem.addprevious(p_img31._element)
    tbl2_elem.getparent().remove(tbl2_elem)
    print("Embedded Figure 3.1 large high-res image.")

    # 3. Insert Figure 3.2 as a large high-res picture
    p_img32 = doc.add_paragraph()
    p_img32.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img32.paragraph_format.space_before = Pt(10)
    p_img32.paragraph_format.space_after = Pt(4)
    r32 = p_img32.add_run()
    r32.add_picture(FIG_3_2, width=Inches(6.5))
    tbls_fig32_elem[0].addprevious(p_img32._element)
    for t_el in tbls_fig32_elem:
        t_el.getparent().remove(t_el)
    print("Embedded Figure 3.2 large high-res image.")

    # 4. Insert Figure 3.3 as a large high-res picture
    p_img33 = doc.add_paragraph()
    p_img33.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img33.paragraph_format.space_before = Pt(10)
    p_img33.paragraph_format.space_after = Pt(4)
    r33 = p_img33.add_run()
    r33.add_picture(FIG_3_3, width=Inches(6.5))
    tbls_fig33_elem[0].addprevious(p_img33._element)
    for t_el in tbls_fig33_elem:
        t_el.getparent().remove(t_el)
    print("Embedded Figure 3.3 large high-res image.")

    # 5. Insert Figure 3.4 as a large high-res picture
    p_img34 = doc.add_paragraph()
    p_img34.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img34.paragraph_format.space_before = Pt(10)
    p_img34.paragraph_format.space_after = Pt(4)
    r34 = p_img34.add_run()
    r34.add_picture(FIG_3_4, width=Inches(6.5))
    tbls_fig34_elem[0].addprevious(p_img34._element)
    for t_el in tbls_fig34_elem:
        t_el.getparent().remove(t_el)
    print("Embedded Figure 3.4 large high-res image.")

    # 6. Rebuild Table 3.1 with Google Colab, Kaggle, Linux VPS, Google Play Store
    new_hw_table = doc.add_table(rows=6, cols=4)
    new_hw_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hw_headers = ["Environment", "Hardware / Platform Specification", "Primary Role in Lifecycle", "Financial Cost / Budget"]
    for idx, text in enumerate(hw_headers):
        c = new_hw_table.rows[0].cells[idx]
        c.text = text
        set_cell_shading(c, "1A5D1A")
        set_cell_borders(c, top="1A5D1A", bottom="1A5D1A", left="1A5D1A", right="1A5D1A")
        set_cell_margins(c, top=100, bottom=100, left=140, right=140)
        p = c.paragraphs[0]
        style_paragraph(p, font_name="Times New Roman", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line_spacing=1.0)
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)

    hw_rows = [
        ("WSL2 — Ubuntu 24.04 (Local)", "Intel/AMD x86_64 CPU, 5.6 GB RAM allocated, Windows Host", "Local backend development, test UI server, sentence segmentation", "Zero cost (Existing hardware)"),
        ("Kaggle GPU Environment", "NVIDIA Tesla T4 GPU (sm_75), 16 GB GDDR6 VRAM, 30 hrs/wk", "Initial QLoRA fine-tuning (v1.0–v1.1); encountered session resets & quotas", "Zero cost (Academic free tier)"),
        ("Google Colaboratory (Colab)", "NVIDIA Tesla T4 GPU (sm_75), 15 GB VRAM, High-RAM host", "Secondary resilient training failover for refined models (v1.2–v1.4) & index build", "Zero cost (Compute runtime tier)"),
        ("Production Linux VPS Server", "Cloud Virtual Private Server (2 vCPU, 4 GB RAM, Ubuntu LTS)", "Hosting production FastAPI backend daemon and Web client portal", "$15.00 / month (Cloud hosting)"),
        ("Mobile Client & Deployment", "Google Play Store Developer Console (Android Mobile APK)", "Public deployment and distribution of native Android mobile app", "$25.00 one-time fee (Google Play Console)")
    ]

    for r_idx, r_data in enumerate(hw_rows, start=1):
        shading = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            c = new_hw_table.rows[r_idx].cells[c_idx]
            c.text = val
            set_cell_shading(c, shading)
            set_cell_borders(c, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC")
            set_cell_margins(c, top=90, bottom=90, left=120, right=120)
            p = c.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.RIGHT if c_idx == 3 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, font_name="Times New Roman", size_pt=9.5, bold=(c_idx == 0), align=align, space_after=1, space_before=1, line_spacing=1.0)

    tbl_hw_elem.addprevious(new_hw_table._element)
    tbl_hw_elem.getparent().remove(tbl_hw_elem)
    print("Replaced Table 3.1 with updated hardware specifications.")

    # 7. Update Table 3.2 (Software)
    tbl_sw = doc.tables[1]
    for r in tbl_sw.rows:
        if "React Native" in r.cells[1].text:
            r.cells[3].text = "Cross-platform mobile UI; compiled for Android OS (iOS excluded due to developer fee constraints)"

    # 8. Clean text paragraphs, down arrows, and captions
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if idx > 40:
            if "Table 3.3 (Architecture):" in txt:
                p.text = ""
            elif txt.startswith("Figure 3.1:"):
                make_caption_p(p, "Figure 3.1: Ruzivo System Architecture — Five-Layer Design")
            elif txt.startswith("Figure 3.2:"):
                make_caption_p(p, "Figure 3.2: Activity Diagram — Chat Request Processing Flow")
            elif txt.startswith("Figure 3.3:"):
                make_caption_p(p, "Figure 3.3: Process Flow — Textbook Corpus & RAG Retrieval Pipeline")
            elif txt.startswith("Figure 3.4:"):
                make_caption_p(p, "Figure 3.4: Process Flow — User Authentication and Session Security")
            elif txt == "Table 3.4: Training Corpus Segment Summary":
                make_caption_p(p, "Table 3.5: Training Corpus Segment Summary")
            elif txt == "Table 3.5: LLM Fine-Tuning Hyperparameters":
                make_caption_p(p, "Table 3.6: LLM Fine-Tuning Hyperparameters")
            elif txt == "Table 3.6: Language Model Training Metrics — All Versions":
                make_caption_p(p, "Table 3.7: Language Model Training Metrics — All Versions")

        if txt in ["↓", "YES path ↓", "─── LOGIN FLOW ───"]:
            p.text = ""

    # Remove empty paragraphs
    for p in doc.paragraphs:
        if not p.text.strip() and not p._element.xpath('.//w:drawing'):
            p_elem = p._element
            if p_elem.getparent() is not None:
                p_elem.getparent().remove(p_elem)

    # 9. Update Section 3.2.1 text (hardware, Kaggle challenges, Google Colab failover, VPS, Play Store, iOS exclusion)
    for p in doc.paragraphs:
        if "The Ruzivo system was developed and deployed across three distinct hardware environments" in p.text:
            p.text = (
                "The Ruzivo system was designed, developed, and provisioned across five distinct computing environments, each fulfilling a specialized operational "
                "role across the development, neural training, failover execution, deployment, and user distribution lifecycle. "
                "The development lifecycle utilized a local Windows Subsystem for Linux 2 (WSL2) instance for rapid prototyping and service integration. "
                "Neural training and vector index generation were initially executed using cloud-based Kaggle GPU environments leveraging NVIDIA Tesla T4 hardware. "
                "However, following severe operational challenges in Kaggle—including unexpected runtime disconnects, strict 30-hour weekly quota exhaustion, "
                "and catastrophic training session restarts that destroyed checkpoints in early model runs—the neural training workflow was migrated to Google Colaboratory (Colab) "
                "as a high-reliability failover computing platform. "
                "For production hosting, the backend REST services and Web client portal are hosted on a dedicated Linux Virtual Private Server (VPS), "
                "budgeted at approximately $15.00 per month. "
                "For mobile distribution, the native client application is targeted exclusively for the Android ecosystem via the Google Play Store, "
                "requiring a one-time Google Play Console developer registration fee of $25.00. "
                "Deployment to Apple's iOS ecosystem was deliberately excluded from the project scope due to prohibitive financial and infrastructural barriers, "
                "specifically the recurring $99.00 annual Apple Developer Program subscription fee and mandatory proprietary macOS hardware requirements. "
                "Table 3.1 details the hardware and platform specifications alongside their respective financial allocations."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

        elif "3.2.1.2 Training Environment — Kaggle GPU Notebooks" in p.text:
            p.text = "3.2.1.2 Training Environments — Kaggle Notebooks and Google Colaboratory (Colab)"
            style_paragraph(p, size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=2)

        elif "Kaggle Notebooks provided free access to T4 GPU instances" in p.text:
            p.text = (
                "Initial model adaptation experiments were conducted on Kaggle Notebooks, which provided access to NVIDIA Tesla T4 GPU instances (sm_75, 16 GB GDDR6 VRAM) "
                "subject to a 30-hour weekly quota. However, Kaggle presented critical operational bottlenecks during long-running supervised fine-tuning. "
                "Specifically, Kaggle's non-interactive execution sessions frequently disconnected unexpectedly without persisting volatile workspace data, "
                "and strict 12-hour session limits triggered mid-training restarts (which directly caused the catastrophic model collapse in version 1.0 at step 3,600). "
                "Furthermore, weekly GPU quota exhaustion halted iterative debugging cycles.\n\n"
                "To resolve these operational constraints and establish computing redundancy, the project incorporated Google Colaboratory (Colab) as a secondary and failover GPU training runtime. "
                "Google Colab provided immediate on-demand NVIDIA T4 GPU instances with 15 GB of video memory and scalable system RAM (up to 25 GB in High-RAM mode). "
                "Colab seamlessly integrated with Google Drive for continuous non-volatile checkpoint backups, shielding training progress against sudden runtime interruptions. "
                "Subsequent stable model iterations (versions 1.2 through 1.4) and dense vector index generation were successfully executed and validated using Google Colaboratory runtimes."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

        elif "3.2.1.3 Storage Environment — Hugging Face Hub" in p.text:
            p.text = "3.2.1.3 Production Hosting Environment — Linux VPS Server"
            style_paragraph(p, size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=2)

        elif "The Hugging Face Hub is a cloud-based platform for hosting and sharing machine learning models" in p.text:
            p.text = (
                "To ensure 24/7 service availability, low-latency API access for mobile and web users, and persistent SQLite database reliability, "
                "the Ruzivo backend daemon is hosted on a cloud-based Linux Virtual Private Server (VPS) provisioned at a projected operational expenditure of $15.00 per month. "
                "A Virtual Private Server provides dedicated virtual CPU cores, allocated memory, and an isolated operating system environment running Ubuntu LTS. "
                "This server environment executes the unified asynchronous Python HTTP service, manages SQLite user authentication and session tables, "
                "and hosts the production web client. Mobile distribution is managed through the Google Play Store for Android devices, leveraging an official developer registration "
                "costing a one-time fee of $25.00. Hugging Face Hub is maintained as the complementary persistent model repository for Git LFS versioned checkpoints."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

        elif "React Native is a JavaScript framework developed by Meta Platforms" in p.text:
            p.text = (
                "React Native is a cross-platform mobile application framework developed by Meta Platforms that compiles JavaScript and React components into native platform widgets. "
                "The Ruzivo client application was constructed using Expo SDK 56, providing a unified codebase for both mobile and web interfaces. "
                "The mobile target was strictly specified for Android devices, supporting deployment via the Google Play Store ($25 developer fee). "
                "Android was prioritized because over 85% of mobile users in Zimbabwe utilize Android-powered smartphones. "
                "iOS compilation and App Store publishing were omitted due to the disproportionate financial overhead ($99/year developer fee) and hardware prerequisites."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

    # 10. Update Section 3.4.2 (Scribd acquisition & 19 books table)
    target_p91 = None
    for p in doc.paragraphs:
        if "The Ruzivo text corpus was assembled from four publicly available" in p.text:
            target_p91 = p
            break

    if target_p91:
        new_text_342 = (
            "The initial text corpus for the Ruzivo project was assembled from four publicly available Shona-language digital sources: "
            "the Google FLEURS Shona speech transcriptions (sn_zw) [20] (3,781 sentences), the digital Shona Bible (31,029 sentences), "
            "the Shona Wikipedia dump (82,707 sentences), and journalistic excerpts from Nehanda Radio (124 sentences). "
            "While this baseline collection yielded 98,993 deduplicated sentences, initial conversational testing revealed critical pedagogical weaknesses. "
            "General web dumps and religious texts lacked rigorous coverage of structured Shona grammatical concepts, such as noun classes (Mipanda yeMazita 1–21), "
            "ideophones (Nyaudzosingwi), traditional proverbs (Tsumo neZvirevo), riddles (Zvirahwe), and secondary-school curriculum literature. "
            "When queried on grammatical definitions, baseline models produced hallucinations or defaulted to English explanations.\n\n"
            "To overcome this domain limitation and ground the system in authentic Zimbabwean school curriculum material, a dedicated primary data collection "
            "and digitization exercise was conducted. A corpus of 19 authoritative Shona educational textbooks, study packs, and linguistic reference guides "
            "was acquired through an educational digital access subscription on Scribd, incurring an acquisition expenditure of $8.00. "
            "These source materials were processed into optical character recognition (OCR) plain text format (.txt), located in the project's data/books_ocr/ directory. "
            "This literature collection comprises 1,739,585 bytes (1.66 MB) of comprehensive pedagogical content, spanning primary school level (Pass Your Grade 7 Shona, "
            "Rodza Pfungwa Grade 5), secondary curriculum (O-Level Shona Study Pack, Form 1–2 Mutauro Notes, Ngatidzidzei ChiShona Form 2, Ngatizivei Mutauro), "
            "and advanced cultural and grammatical treatises (Mipanda yeMazita, Nyaudzosingwi, Tsumo, Zvindori, and A-Level Literature Analysis). "
            "Table 3.4 outlines the complete inventory of the 19 educational text sources assembled for this dataset."
        )
        target_p91.text = new_text_342
        style_paragraph(target_p91, size_pt=12, bold=False, line_spacing=1.5)

        tbl_caption = doc.add_paragraph()
        make_caption_p(tbl_caption, "Table 3.4: Inventory of Curated Shona Educational Books and Reference Texts")
        target_p91._element.addnext(tbl_caption._element)

        table_books = doc.add_table(rows=20, cols=4)
        table_books.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["File Name (.txt)", "Domain / Curriculum Level", "File Size", "Pedagogical Focus & Acquisition Source"]
        hdr_cells = table_books.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_shading(hdr_cells[idx], "1A5D1A")
            p = hdr_cells[idx].paragraphs[0]
            style_paragraph(p, font_name="Times New Roman", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line_spacing=1.0)
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)

        books_data = [
            ("529832404-Shona-Module-ChiShona-Vol-1.txt", "Secondary / Tertiary", "496.2 KB", "Comprehensive phonology, morphology & syntax (Scribd OCR)"),
            ("474646692-O-LEVEL-SHONA-STUDY-PACK.txt", "ZIMSEC O-Level", "330.7 KB", "Grammar exercises, literature, comprehension (Scribd OCR)"),
            ("616354923-Shona-Peqs-and-Answers.txt", "ZIMSEC Exam Prep", "204.4 KB", "Exam past questions, model Shona solutions (Scribd OCR)"),
            ("486365494-Pass-Your-Grade-7-Shona.txt", "Primary Grade 7", "172.1 KB", "Foundational grammar, vocabulary, proverbs (Scribd OCR)"),
            ("620864876-Rodza-Pfungwa-G5-Sample.txt", "Primary Grade 5", "96.5 KB", "Early comprehension, sentence construction (Scribd OCR)"),
            ("837886802-NGATIZIVEI-MUTAURO.txt", "Secondary Mutauro", "80.9 KB", "Dialectal variations, word structures, idioms (Scribd OCR)"),
            ("570127832-Ngatidzidzei-ChiShona-F2.txt", "Secondary Form 2", "76.6 KB", "Intermediate linguistic exercises & syntax (Scribd OCR)"),
            ("746517921-SHONA-NOTES-2022.txt", "Secondary Revision", "49.2 KB", "Curriculum study summaries & language rules (Scribd OCR)"),
            ("649864977-MIPANDA-YEMAZITA-1.txt", "Linguistic Grammar", "46.6 KB", "Exhaustive noun class system (Classes 1–21) (Scribd OCR)"),
            ("648101854-Nyaudzosingwi.txt", "Linguistic Grammar", "42.5 KB", "Ideophone taxonomy, sound symbolism & usage (Scribd OCR)"),
            ("995907723-SirD-NHETEMBO-unlocked.txt", "Literature / Poetry", "32.9 KB", "Traditional poetry, rhythm & poetic devices (Scribd OCR)"),
            ("981493047-Shona-Lit-Ongororo-A-Level.txt", "ZIMSEC A-Level", "31.0 KB", "Literary analysis, cultural critique & themes (Scribd OCR)"),
            ("576008136-Tsumo.txt", "Cultural Heritage", "24.6 KB", "Categorised traditional proverbs & contexts (Scribd OCR)"),
            ("431120508-TSUMO.txt", "Cultural Heritage", "17.9 KB", "Proverbs with literal & metaphorical meanings (Scribd OCR)"),
            ("843345121-Form-1-2-mutauro-notes.txt", "Secondary Form 1–2", "15.3 KB", "Core parts of speech, noun prefixes, concords (Scribd OCR)"),
            ("812809917-Literature-Questions-2023.txt", "Literature Exam Prep", "12.3 KB", "Modern and traditional literary essay prompts (Scribd OCR)"),
            ("853697528-MHENENGURO-NOTSI-1.txt", "Literary Analysis", "5.6 KB", "Advanced prose critique and thematic evaluation (Scribd OCR)"),
            ("885600765-ZVINDORI-1.txt", "Folklore / Riddles", "2.7 KB", "Cultural riddles, word puzzles & oral traditions (Scribd OCR)"),
            ("Zvivakashure neMhando dzeZvipauro.txt", "Morphology Guide", "1.5 KB", "Noun prefixes, adjective stems & concord rules (Scribd OCR)")
        ]

        for row_idx, row_data in enumerate(books_data, start=1):
            row_cells = table_books.rows[row_idx].cells
            shading_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, text in enumerate(row_data):
                row_cells[col_idx].text = text
                set_cell_shading(row_cells[col_idx], shading_col)
                p = row_cells[col_idx].paragraphs[0]
                align = WD_ALIGN_PARAGRAPH.RIGHT if col_idx == 2 else (WD_ALIGN_PARAGRAPH.LEFT if col_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER)
                style_paragraph(p, font_name="Times New Roman", size_pt=9.5, bold=(col_idx == 0), align=align, space_after=1, space_before=1, line_spacing=1.0)

        tbl_caption._element.addnext(table_books._element)

        p_pipe = doc.add_paragraph()
        p_pipe_text = (
            "The raw educational texts acquired from Scribd underwent a structured multi-stage data curation pipeline. "
            "First, character encoding inconsistencies and OCR noise (e.g., misrecognized diacritics, broken ligatures, "
            "and stray scanning artifacts) were stripped via targeted regular expression transforms. "
            "Second, running headers, footers, publisher page numbering, and repetitive exercise borders were systematically excised. "
            "Third, the sanitized literature was processed through a custom sentence segmentation tokenizer tailored to Shona orthography, "
            "handling honorific abbreviations and dialogue punctuation without erroneous mid-sentence truncation. "
            "Fourth, sentences were filtered against strict length constraints (retaining units between 4 and 45 words) and evaluated "
            "against the corpus-vocabulary-subtraction filter to verify zero residual English contamination. "
            "This automated refinement reduced the 1.66 MB raw textbook corpus into 14,201 verified, pedagogically dense Shona sentences "
            "persisted in UTF-8 JSON format at data/rag_knowledge_base/sentences.json (919 KB). "
            "These 14,201 sentences serve as the authoritative factual context repository for the system's online Retrieval-Augmented Generation pipeline."
        )
        p_pipe.text = p_pipe_text
        style_paragraph(p_pipe, size_pt=12, bold=False, line_spacing=1.5)
        table_books._element.addnext(p_pipe._element)

    # 11. Update Section 3.5.1, 3.6.1 & Summary
    for p in doc.paragraphs:
        if "A lightweight, rule-based language detection algorithm was implemented" in p.text:
            p.text = (
                "A robust multi-tier language detection and filtering algorithm was engineered to enforce strict ChiShona-only communication. "
                "Because off-the-shelf statistical classifiers (such as fastText or langdetect) exhibit high error rates on short, morphologically rich "
                "Bantu phrases and impose prohibitive inference latency overheads, Ruzivo deploys a hybrid deterministic lexical gatekeeper (is_strictly_non_shona). "
                "The gatekeeper extracts alphabetic tokens and compares them against a curated lexicon of high-frequency English functional words. "
                "A query is intercepted and rejected if: (1) it contains 3 or fewer words matching common foreign salutations or queries (e.g., 'hello', 'hi', 'morning', 'help', 'translate'), "
                "or (2) non-Shona words constitute 25% or more of the total word count. "
                "Non-compliant inputs are immediately rejected at the gateway level with a polite, authentic ChiShona refusal ('Chirongwa chino chinongoshandisa mutauro weChiShona chete'), "
                "completely bypassing neural execution and eliminating unauthorized English processing."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

        elif "The Ruzivo language model was implemented as a QLoRA-adapted instance" in p.text:
            p.text = (
                "The Ruzivo conversational intelligence layer employs a dual-tier neural architecture: a localized fine-tuned Qwen2.5-1.5B [12] model "
                "adapted via QLoRA for standalone offline environments, paired with an advanced multi-tier neural generative service for complex conversational reasoning. "
                "To overcome earlier empirical failures—where the model generated excessive boilerplate paragraphs, boasted about its origin without prompting, "
                "or repeated historical creator narratives—a strict cognitive instruction framework was formalized within the generation pipeline. "
                "The core engine operates under six mandatory behavioral imperatives: "
                "(1) Proportionate Response Length: answers strictly scale to query complexity; simple greetings ('Mangwanani', 'Masikati') receive warm, concise 1-to-2 sentence replies rather than multi-paragraph essays; "
                "(2) Origin Shielding: the model is strictly forbidden from reciting its development history or institutional affiliations unless explicitly asked ('Wakasikwa naani?'); "
                "(3) Exclusive ChiShona Synthesis: output tokens must be 100% natural, grammatically correct Shona; "
                "(4) Natural Idiomatic Register: proverbs (tsumo) and idioms (madimikira) are introduced only when contextually relevant rather than force-injected into every turn; "
                "(5) Factual Textbook Grounding: educational queries strictly defer to verified evidence extracted from the 14,201-sentence textbook knowledge base; and "
                "(6) Hallucination Suppression: non-existent concepts are refused with an explicit 'Handizivi' rather than speculative generation."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

        elif "This chapter described the complete methodology of the Ruzivo project" in p.text:
            p.text = (
                "This chapter presented the comprehensive methodology of the Ruzivo conversational AI platform within the Design Science Research (DSR) framework [5]. "
                "Hardware and operational requirements were established across local WSL2 development, Kaggle T4 GPU fine-tuning, Google Colaboratory failover runtime environments, "
                "cloud Linux VPS server hosting ($15/month), and native Android mobile application distribution via the Google Play Store ($25 developer fee), while omitting costly iOS deployment. "
                "The software architecture established theoretical and operational foundations for PyTorch [23], FastAPI [24], SQLite relational authentication, "
                "FAISS vector indexing [15], and React Native Expo client modules. "
                "Primary data acquisition was documented across two complementary vectors: community crowd-sourcing via Google Forms and the acquisition of 19 authoritative "
                "Shona educational textbooks via Scribd ($8 subscription), preprocessed from 1.66 MB of raw text into 14,201 verified pedagogical sentences. "
                "The five-layer system architecture was illustrated in Figure 3.1, while activity and process flows for chat processing, RAG context retrieval, "
                "and user authentication security were formalized in Figures 3.2 through 3.4. "
                "Finally, algorithmic designs for strict ChiShona lexical gatekeeping, proportionate length generation, origin shielding, and multi-turn in-context memory "
                "were articulated, establishing a robust, culturally authentic conversational agent."
            )
            style_paragraph(p, size_pt=12, bold=False, line_spacing=1.5)

    # 12. Front-matter TOC update
    for idx, p in enumerate(doc.paragraphs[:40]):
        if "Table 3.4: Training Corpus Segment Summary" in p.text:
            p.text = "Table 3.4: Inventory of Curated Shona Educational Books and Reference Texts\t13"
            p_new = doc.add_paragraph()
            p_new.text = "Table 3.5: Training Corpus Segment Summary\t16"
            style_paragraph(p_new, font_name="Times New Roman", size_pt=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.0)
            p._element.addnext(p_new._element)
        elif "Table 3.5: LLM Fine-Tuning Hyperparameters" in p.text:
            p.text = "Table 3.6: LLM Fine-Tuning Hyperparameters\t25"
        elif "Table 3.6: Language Model Training Metrics" in p.text:
            p.text = "Table 3.7: Language Model Training Metrics — All Versions\t31"

    # 13. Format headings
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "Chapter 3: Methodology":
            style_paragraph(p, font_name="Times New Roman", size_pt=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=12)
        elif any(txt.startswith(prefix) for prefix in ["3.1 ", "3.2 ", "3.3 ", "3.4 ", "3.5 ", "3.6 ", "3.7 ", "3.8 ", "References", "List of Tables", "List of Figures", "Table of Contents"]):
            style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4)
        elif any(txt.startswith(f"3.{i}.{j}") for i in range(1, 9) for j in range(1, 15)):
            style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=2)

    # Save to main target output (or fallback if file is open in Word)
    try:
        doc.save(OUT_DOCX)
        print(f"Successfully generated Chapter 3 with large, crystal-clear publication diagrams to: {OUT_DOCX}")
    except PermissionError:
        fallback_path = r"C:\Users\recal\Downloads\Ruzivo_Chapter3_Complete_Updated.docx"
        doc.save(fallback_path)
        print(f"Note: '{OUT_DOCX}' is currently open in Microsoft Word.")
        print(f"Successfully saved updated document to: {fallback_path}")

if __name__ == "__main__":
    update_chapter_3_with_large_diagrams()
